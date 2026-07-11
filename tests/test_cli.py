import json
from pathlib import Path

import pytest
from conftest import FakeDetector

from deident import cli


def test_init_writes_example_config(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    assert cli.main(["init"]) == 0
    content = (tmp_path / "project.yaml").read_text()
    assert "known_names:" in content
    assert "ollama_model:" in content
    # refuses to clobber without --force
    assert cli.main(["init"]) == 1
    assert cli.main(["init", "--force"]) == 0


def test_run_skip_llm_batch(project_dir, capsys):
    rc = cli.main(["run", "--skip-llm", "transcripts"])
    assert rc == 0
    out = capsys.readouterr().out

    cleaned = project_dir / "cleaned"
    assert (cleaned / "synthetic_focus_group_deid.txt").exists()
    assert (cleaned / "synthetic_followup_deid.vtt").exists()
    assert (project_dir / "review_report.md").exists()

    key = json.loads((project_dir / "deident_key.json").read_text())
    assert key["categories"]["name"]["counter"] == 3
    assert "Done: 2 file(s)" in out
    assert "review report" in out

    report = (project_dir / "review_report.md").read_text()
    assert "deterministic only (--skip-llm)" in report
    assert "synthetic_focus_group.txt" in report


def test_run_single_docx_emitted_as_txt(project_dir):
    import docx

    doc = docx.Document()
    doc.add_paragraph("MODERATOR: Welcome to St. Aurelia Hospital.")
    doc.add_paragraph("MARISOL: Marisol's team covers B4.")
    doc_path = project_dir / "transcripts" / "notes.docx"
    doc.save(str(doc_path))

    assert cli.main(["run", "--skip-llm", str(doc_path)]) == 0
    out_path = project_dir / "cleaned" / "notes_deid.txt"
    cleaned = out_path.read_text()
    assert "Site A" in cleaned
    assert "Participant 01's team covers Unit 1" in cleaned


def test_run_full_pipeline_with_mocked_ollama(project_dir, monkeypatch, capsys):
    detector = FakeDetector(
        if_present=[
            {"text": "Rosewood Clinic", "category": "site", "reason": "clinic"}
        ]
    )
    (project_dir / "transcripts" / "extra.txt").write_text(
        "SPEAKER: They transferred her to Rosewood Clinic overnight.\n"
    )
    monkeypatch.setattr(cli, "OllamaClient", lambda url, model: detector)

    assert cli.main(["run", "transcripts"]) == 0
    assert detector.calls > 0
    cleaned = (project_dir / "cleaned" / "extra_deid.txt").read_text()
    assert "Rosewood" not in cleaned and "Site C" in cleaned
    report = (project_dir / "review_report.md").read_text()
    assert "Rosewood Clinic" in report
    assert "AUTO-REPLACED with 'Site C'" in report


def test_run_missing_config_fails_clearly(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    (tmp_path / "t.txt").write_text("hello")
    assert cli.main(["run", "--skip-llm", "t.txt"]) == 2
    assert "deident init" in capsys.readouterr().err


def test_run_warns_when_key_file_missing_on_rerun(project_dir, capsys):
    assert cli.main(["run", "--skip-llm", "transcripts"]) == 0
    capsys.readouterr()
    (project_dir / "deident_key.json").unlink()
    assert cli.main(["run", "--skip-llm", "transcripts"]) == 0
    err = capsys.readouterr().err
    assert "WARNING" in err and "NOT match earlier runs" in err


def test_status_counts_only(project_dir, capsys):
    assert cli.main(["run", "--skip-llm", "transcripts"]) == 0
    capsys.readouterr()
    assert cli.main(["status"]) == 0
    out = capsys.readouterr().out
    assert "names mapped:  3" in out
    assert "sites mapped:  2" in out
    # actual identifiers never appear in status output
    assert "Marisol" not in out and "Aurelia" not in out


def test_status_without_key_file(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    assert cli.main(["status"]) == 1
    assert "nothing mapped yet" in capsys.readouterr().out


def test_nonlocal_ollama_url_rejected(project_dir, capsys):
    cfg = (project_dir / "project.yaml").read_text()
    (project_dir / "project.yaml").write_text(
        cfg + "ollama_url: https://api.example.com\n"
    )
    assert cli.main(["run", "--skip-llm", "transcripts"]) == 2
    assert "localhost" in capsys.readouterr().err
